#!/usr/bin/env python3
"""Builds portafolio-cms from the original portafolio.html.

  python3 build.py [path/to/portafolio.html]

Outputs
  index.html                          site + hidden admin (single file, deploy anywhere)
  database/hipernova-portafolio-db.xlsx   seed for the Google Sheet (import it into Drive)
  database/seed.json                  the same content as JSON (reference)

Needs `openpyxl` only for the .xlsx (pip install openpyxl).
"""
import json, re, sys, pathlib, importlib.util

HERE = pathlib.Path(__file__).parent
SRC = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "source" / "portafolio.html"
lines = SRC.read_text(encoding="utf8").split("\n")

# ---- 1. pull the embedded DATA out of the original -------------------------
i = next(n for n, l in enumerate(lines) if l.startswith("const DATA = "))
raw = lines[i][len("const DATA = "):].rstrip().rstrip(";")
data = json.loads(raw)
for p in data["projects"]:
    p["industry"] = [x.replace("Goverment", "Government") for x in p["industry"]]   # typo in the original data
    p["status"] = "published"
    p.setdefault("cover", "")
    p["link"] = p.get("link") or ""
(HERE / "database").mkdir(exist_ok=True)
(HERE / "database" / "seed.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf8")

_spec = importlib.util.spec_from_file_location("hn_translations", pathlib.Path(__file__).parent / "src" / "translations.py")
TRN = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(TRN)
assert len(TRN.CASES) == len(data["projects"]), "CASES must match the number of projects"
_cnt, _last = {}, {}
for _p in data["projects"]:
    for _c in _p["client"]:
        _cnt[_c] = _cnt.get(_c, 0) + 1; _last[_c] = max(_last.get(_c, 0), _p.get("year") or 0)
data["config"]["clients"] = sorted(_cnt, key=lambda c: (-_cnt[c], -_last[c], c))   # priority: #projects, then recency
data["es"] = TRN.ES
data["i18n"] = {"es": TRN.DICT}
BODIES_ES = json.loads((HERE / "src" / "case_bodies_es.json").read_text(encoding="utf8"))  # {slug: html_es}
for _p, (_t, _s) in zip(data["projects"], TRN.CASES):
    _p["es"] = {"title": _t, "sub": _s}
    if BODIES_ES.get(_p["slug"]):
        assert len(BODIES_ES[_p["slug"]]) <= 49000, "html_es too long for a Sheets cell: " + _p["slug"]
        _p["es"]["html"] = BODIES_ES[_p["slug"]]

HERO_PHRASES = ["see what we couldn't see before.", "change course.", "question the rules we were given.", "do things differently.",
  "turn limits into possibilities.", "imagine other futures.", "let something new emerge.", "discover new possibilities.",
  "shift the way they look at things.", "break with the predictable.", "reimagine what already exists.", "create from another place.",
  "change the context.", "redesign the conditions.", "open new paths.", "change the questions.", "challenge what was taken for granted.",
  "inhabit new scenarios.", "start new conversations.", "expand what seemed possible.", "connect what was separate.",
  "create new ways of relating.", "give shape to what doesn't exist yet.", "turn uncertainty into possibility.",
  "try out other futures.", "make the invisible visible."]
data["pages"]["home"]["hero"] = "I help people, brands and organizations to"
data["pages"]["home"]["phrases"] = HERO_PHRASES
(HERE / "database" / "seed.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf8")

# ---- 2. rewire the site script: data comes from Sheets (or cache / seed) ----
lines[i] = ""   # no content is embedded: everything comes from the Sheet (or database/seed.json in demo mode)
j = next(n for n, l in enumerate(lines) if l.startswith("const P = DATA.projects"))
I18N_JS = (HERE / "src" / "i18n.js").read_text(encoding="utf8")
lines[j] = I18N_JS + """
let DATA, P, C, PG;
function setData(d){
  const dv = u => (u || "").replace(/https:\/\/drive\.google\.com\/thumbnail\?id=([\w-]+)[^"'\s]*/g, "https://lh3.googleusercontent.com/d/$1=w1600");
  const fx = h => dv(h || "").replace(/<img (?![^>]*referrerpolicy)/g, '<img referrerpolicy="no-referrer" ');
  d.projects.forEach(p => { p.industry = (p.industry || []).map(x => x.replace(/^Goverment/, "Government")); p.cover = dv(p.cover); p.html = fx(p.html); if (p.es && p.es.html) p.es.html = fx(p.es.html); });
  DATA = d; const adm = !!window.ADMIN_ON, v = i18nLocalize(d);
  P = v.projects.filter(p => adm || p.status !== "draft"); C = v.config; PG = v.pages;
}
/* No content lives in this file. First paint uses the last copy cached in the browser; the Sheet refreshes it. */
window.HN_READY = false;
function initialData(){
  try { const s = localStorage.getItem(HN_CONFIG.API_URL ? "hn_cache" : "hn_demo_data"); if (s) { const d = JSON.parse(s); d.projects.forEach(p => p.status = p.status || "published"); return d; } } catch(e){}
  return null;
}
const _d0 = initialData(); if (_d0) { setData(_d0); window.HN_READY = true; }
const loadingView = () => `<section class="phead s-energy"><div class="wrap"><span class="t-meta" data-nt>Loading…</span></div></section>`;
/* clients marquee: rows of names, most relevant first (C.clients order, then #projects / recency) */
function clientRows(n){
  const cnt = {}, last = {};
  P.forEach(p => p.client.forEach(c => { cnt[c] = (cnt[c] || 0) + 1; last[c] = Math.max(last[c] || 0, p.year || 0); }));
  const auto = Object.keys(cnt).sort((a, b) => cnt[b] - cnt[a] || last[b] - last[a] || a.localeCompare(b));
  const pri = (C.clients || []).filter(c => cnt[c]);
  const order = pri.concat(auto.filter(c => !pri.includes(c)));
  const rows = Array.from({length: n}, () => []); order.forEach((c, i) => rows[i % n].push(c));
  return rows.filter(r => r.length);
}
const MQ_DUR = [52, 64, 46, 70, 58];
const MQ_STAR = '<svg class="mq-s" viewBox="0 0 100 100" aria-hidden="true"><path d="M50 2 C54 38 62 46 98 50 C62 54 54 62 50 98 C46 62 38 54 2 50 C38 46 46 38 50 2Z" fill="currentColor"/></svg>';
function clientsMarquee(){
  return clientRows(5).map((r, i) => `<div class="mq-row${i % 2 ? " out" : ""}" style="--d:${MQ_DUR[i % 5]}s;--dir:${i % 2 ? "reverse" : "normal"}"><div class="mq-track">${[0, 1].map(k => `<div class="mq-set"${k ? ' aria-hidden="true"' : ""}>${r.map(c => `<span class="mq-i">${esc(c)}</span>${MQ_STAR}`).join("")}</div>`).join("")}</div></div>`).join("");
}"""
html = "\n".join(lines)

HERO_NEW = """<section class="hero-d s-matter">
    <canvas aria-hidden="true"></canvas>
    <div class="hero-in wrap">
      <p class="t-meta hero-top">${C.disciplines.map(esc).join(" · ")}</p>
      <h1 class="hero-h1"><span class="hero-l">${esc(PG.home.hero || "")}</span><span class="hero-ph" data-phrase></span></h1>
      <div class="hero-bot">
        <div class="btns">${btn("#/work", "See the work", "accent lg")}${btn("#/approach", "How I work", "outline lg")}</div>
        <span class="hero-tag">People · Systems · Futures</span>
      </div>
    </div>
  </section>"""
SIG_OLD = re.search(r'<div><span class="hl">01 — Signal / Field</span>.*?</div>\s*', html, re.S)
assert SIG_OLD, "signal block not found"
html = html[:SIG_OLD.start()] + html[SIG_OLD.end():]
m = re.search(r'<section class="hero s-energy">.*?</section>', html, re.S)
assert m, "old hero not found"
html = html[:m.start()] + HERO_NEW + html[m.end():]
RT_OLD = 'if (r === "work" && !slug) workBind();'
assert RT_OLD in html
RO = 'app.innerHTML = r === "" ? home()'
assert RO in html, "route anchor"
html = html.replace(RO, 'app.innerHTML = !window.HN_READY ? loadingView() : r === "" ? home()', 1)
assert 'if (r === "work" && !slug) workBind();' in html
RT_NEW = 'if (window.HN_READY && r === "work" && !slug) workBind();'
html = html.replace(RT_OLD, RT_NEW, 1)
RT_OLD = RT_NEW   # the hook chain below is appended after this statement
html = html.replace(RT_OLD, RT_OLD + ' window.initHero && initHero();', 1)
assert 'window.initHero && initHero();' in html

METHOD = """<section class="method-s s-matter">
    <div class="wrap method-in">
      <div class="method-l">
        <span class="t-meta method-k">01 — Method · System</span>
        <h2 class="method-h">A cycle of four movements.</h2>
        <div class="method-card">
          <span class="t-meta method-n hl" data-m="num">01 / 04</span>
          <p class="method-name" data-m="name">OBSERVE</p>
          <p class="method-desc" data-m="desc"></p>
        </div>
      </div>
      <div class="method-r">
        <canvas aria-hidden="true"></canvas>
        <div class="method-core" aria-hidden="true"><span>CYCLE</span><span data-m="mode">SIGNAL</span></div>
        <div data-m="nodes"></div>
      </div>
    </div>
  </section>
  """
APP_OLD = '<section class="s-space"><div class="wrap">\n    ${sec(1, "Six principles"'
assert APP_OLD in html, "approach anchor not found"
html = html.replace(APP_OLD, METHOD + '<section class="s-space"><div class="wrap">\n    ${sec(2, "Six principles"', 1)
html = html.replace(' window.initHero && initHero();', ' window.initHero && initHero(); window.initMethod && initMethod();')

TH_OLD = '<div class="btns">${btn(C.substack, "Read on Substack", "accent")}</div>`)}'
assert TH_OLD in html, "thinking anchor not found"
html = html.replace(TH_OLD, '<div class="btns">${btn(C.substack || "https://hipernova.substack.com", "Read on Substack", "accent")}</div><div class="posts" id="posts"></div>`)}', 1)
assert ' window.initMethod && initMethod();' in html
html = html.replace(' window.initMethod && initMethod();', ' window.initMethod && initMethod(); window.initPosts && initPosts();')

HDR_OLD = '</nav>\n  </div>\n</header>'
assert HDR_OLD in html, "header anchor not found"
html = html.replace(HDR_OLD, '</nav>\n    <div class="lang" role="group" aria-label="Language"><button type="button" data-l="en">EN</button><button type="button" data-l="es">ES</button></div>\n  </div>\n</header>', 1)
html = html.replace(' window.initPosts && initPosts();', ' window.initPosts && initPosts(); i18nDom();', 1)
html = html.replace('addEventListener("hashchange", route);', 'addEventListener("hashchange", route);\ndocument.querySelector(".lang").addEventListener("click", e => { const b = e.target.closest("button"); if (b) setLang(b.dataset.l); });\ndocument.querySelectorAll(".lang button").forEach(b => { b.classList.toggle("on", b.dataset.l === LANG); b.setAttribute("aria-pressed", b.dataset.l === LANG); });', 1)

MEAN = """</div>
  <section class="mean-s s-matter">
    <div class="mean-st">
      <canvas aria-hidden="true"></canvas>
      <div class="mean-in wrap">
        <div class="mean-l">
          <span class="t-meta mean-k">02 — What is HIPERNOVA?</span>
          <p class="mean-p">A hypernova is a concentration of energy capable of altering what surrounds it. I'm not interested in the astronomical phenomenon, but in its behavior.</p>
          <div class="mean-card"><span class="t-meta hl" data-mm="count">Signal — 00 / 06</span><p class="mean-t" data-mm="title"></p><p class="mean-d" data-mm="text"></p></div>
        </div>
        <ul class="mean-words" data-mm="words"></ul>
      </div>
    </div>
  </section>
  <section class="s-matter mean-band"><div class="wrap"><p><span class="mean-dim">My intention isn't to produce change.</span> <span class="hl">It's to alter trajectories.</span></p></div></section>
  <div class="wrap">
    ${sec(3, "Latest writing", `<h2>From the Substack.</h2><div class="posts" id="home-posts"></div><div class="btns">${btn("#/thinking", "All writing")}</div>`)}
    """
H_OLD = '${sec(2, "Practice"'
assert H_OLD in html, "home practice anchor"
html = html.replace(H_OLD, MEAN + '${sec(4, "Practice"', 1)
for o, n in [('${sec(3, "Principles"', '${sec(5, "Principles"'), ('${sec(4, "Clients & partners"', '${sec(6, "Clients & partners"'), ('<span class="t-meta">05 — Next</span>', '<span class="t-meta">07 — Next</span>')]:
    assert o in html, o
    html = html.replace(o, n, 1)
html = html.replace(' window.initMethod && initMethod();', ' window.initMethod && initMethod(); window.initMeaning && initMeaning();', 1)

import re as _re
_m = _re.search(r'\$\{sec\(\d, "Clients & partners".*?\)\}\n', html, _re.S)
assert _m, "clients section not found"
MQ = """</div>
  <section class="clients-s s-space">
    <div class="wrap"><span class="t-meta">06 — Clients & partners</span><h2 class="mq-h">Who I've worked with.</h2></div>
    <div class="mq" aria-label="Clients and partners">${clientsMarquee()}</div>
  </section>
  <div class="wrap">
  """
html = html[:_m.start()] + MQ + html[_m.end():]

# cover photo (duotone) in cards and case band; texture stays as fallback
CARD_OLD = '<canvas class="ht" data-m="${MOTIF[p.texture] || "dots"}" data-seed="${P.indexOf(p) + 3}"></canvas>'
CARD_NEW = '${p.cover ? `<img class="duo" referrerpolicy="no-referrer" src="${esc(p.cover)}" alt="" loading="lazy" data-k="card" data-t="${esc(p.texture)}" data-s="${P.indexOf(p) + 3}">` : `' + CARD_OLD + '`}'
BAND_OLD = '<canvas data-hn-texture="${p.texture}" data-tone="ink" data-seed="${i + 7}"></canvas>'
BAND_NEW = '${p.cover ? `<img class="duo" referrerpolicy="no-referrer" src="${esc(p.cover)}" alt="" data-k="band" data-t="${esc(p.texture)}" data-s="${i + 7}">` : `' + BAND_OLD + '`}'
assert CARD_OLD in html and BAND_OLD in html
html = html.replace(CARD_OLD, CARD_NEW).replace(BAND_OLD, BAND_NEW)
DUO = """<svg width="0" height="0" style="position:absolute" aria-hidden="true"><defs><filter id="hn-duo" color-interpolation-filters="sRGB">
<feColorMatrix type="matrix" values="0.33 0.5 0.17 0 0  0.33 0.5 0.17 0 0  0.33 0.5 0.17 0 0  0 0 0 1 0"/>
<feComponentTransfer><feFuncR type="table" tableValues="0.133 0.992"/><feFuncG type="table" tableValues="0.133 0.933"/><feFuncB type="table" tableValues="0.133 0.239"/></feComponentTransfer></filter></defs></svg>
<style>.duo{filter:url(#hn-duo);width:100%;height:100%;object-fit:cover;display:block}.band .duo{position:absolute;inset:0}.card .viz .duo{position:absolute;inset:0}.duo-ht{filter:none!important}</style>
"""
html = html.replace("<body class=\"s-space\">", "<body class=\"s-space\">\n" + DUO, 1)

config = """<script>
/* ===== HIPERNOVA CMS — connection ==========================================
   API_URL : the "Web app" URL of your Google Apps Script (see README).
             Leave empty = DEMO mode (edits are saved in this browser only).
   COVER_STYLE : "mix" (duotone + halftone dots) | "halftone" | "smooth" (duotone only)
   Admin   : click the logo 5 times. The real password lives in Apps Script
             (Script properties → ADMIN_PASSWORD); DEMO_PASSWORD only applies in demo mode. */
window.HN_CONFIG = { API_URL: "https://script.google.com/macros/s/AKfycbzAY0I4SW7_5R5KyOt1Qx7cIuCrq3cLd71XzSAPL7o1IMH0mY0EWBHNY4FmK3J6onAc8w/exec", DEMO_PASSWORD: "hipernova", CLICKS: 5, COVER_STYLE: "mix" };
</script>
"""
css = (HERE / "src" / "admin.css").read_text(encoding="utf8") + """
/* ---------- home hero (dark, particle field) ---------- */
.hero-d{position:relative;min-height:calc(100vh - 72px);display:flex;overflow:hidden}
.hero-d canvas{position:absolute;inset:0;width:100%;height:100%;display:block}
.hero-in{position:relative;z-index:1;pointer-events:none;display:grid;grid-template-rows:auto 1fr auto;gap:40px;padding-top:clamp(28px,5vw,64px);padding-bottom:clamp(32px,5vw,56px);width:100%}
.hero-in .btns{pointer-events:auto;margin:0}
.hero-top{margin:0;color:#D7D7D7}
.hero-h1{align-self:center;margin:0;max-width:19ch;font:700 clamp(36px,5.4vw,92px)/1.02 var(--font-display);letter-spacing:-.035em;color:#F4F4EE}
.hero-l{display:block;text-wrap:balance}
.hero-ph{display:grid;position:relative;color:var(--hn-yellow);margin-top:.06em}
.hero-bot{display:flex;justify-content:space-between;align-items:flex-end;gap:32px;flex-wrap:wrap}
.hero-tag{font:500 clamp(12px,1.1vw,14px)/1 var(--font-mono);letter-spacing:.2em;text-transform:uppercase;color:#D7D7D7}
.hero-d .hn-btn.accent:hover{background:#F4F4EE;color:#222}
.hero-d .hn-btn.outline{border-color:#F4F4EE;color:#F4F4EE}
.hero-d .hn-btn.outline:hover{background:#F4F4EE;color:#222}
.signal{border-top:1px solid rgba(244,244,238,.18)}
/* Spanish words are longer: keep big page titles clear of the corner texture */
html[lang="es"] .phead h1{font-size:clamp(40px,7.2vw,140px);overflow-wrap:break-word}
html[lang="es"] .phead:has(.corner) h1{max-width:64%}
.lang{display:flex;border:1px solid var(--hn-deep-space);flex:none}
.lang button{font:500 11px/1 var(--font-mono);letter-spacing:var(--tracking-meta);padding:9px 10px;border:0;background:transparent;color:inherit;cursor:pointer}
.lang button.on{background:var(--hn-deep-space);color:var(--hn-yellow)}
.lang button:not(.on):hover{background:rgba(34,34,34,.12)}
.no-es{font:400 14px/1.5 var(--font-mono);opacity:.7;border-left:3px solid var(--hn-yellow);padding-left:12px}
.posts{margin-top:56px}.posts-n{margin:0 0 16px}
.signal-in>div:nth-child(1){grid-column:1/-1;display:flex;gap:28px 48px;flex-wrap:wrap}
@media (pointer:fine){html.hn-cursor,html.hn-cursor *{cursor:none!important}
html.hn-cursor #adm,html.hn-cursor #adm *,html.hn-cursor .mdl-bg,html.hn-cursor .mdl-bg *{cursor:auto!important}
html.hn-cursor #adm button,html.hn-cursor #adm select,html.hn-cursor #adm label,html.hn-cursor .mdl-bg button{cursor:pointer!important}
html.hn-cursor #adm input,html.hn-cursor #adm textarea,html.hn-cursor #adm [contenteditable],html.hn-cursor .mdl-bg input{cursor:text!important}}
/* ---------- home › clients marquee ---------- */
.clients-s{padding:clamp(72px,9vw,128px) 0 clamp(56px,7vw,96px);border-top:var(--border-bold);overflow:hidden}
.mq-h{margin:20px 0 clamp(36px,5vw,64px);font:700 clamp(36px,4.6vw,56px)/.98 var(--font-display);letter-spacing:-.03em;text-transform:uppercase}
.mq{display:grid;gap:clamp(8px,1.2vw,18px)}
.mq-row{overflow:hidden;white-space:nowrap}
.mq-track{display:flex;width:max-content;animation:mq var(--d,50s) linear infinite;animation-direction:var(--dir,normal)}
.mq-row:hover .mq-track{animation-play-state:paused}
.mq-set{display:flex;align-items:center;flex:none;gap:clamp(24px,3vw,48px);padding-right:clamp(24px,3vw,48px)}
.mq-i{font:700 clamp(40px,7.2vw,112px)/1.05 var(--font-display);letter-spacing:-.04em;text-transform:uppercase;padding:0 .08em;transition:background var(--dur-fast) var(--ease-field)}
.mq-row.out .mq-i{color:transparent;-webkit-text-stroke:2px var(--hn-deep-space)}
.mq-i:hover{background:var(--hn-yellow);color:var(--hn-deep-space);-webkit-text-stroke:0}
.mq-s{width:clamp(18px,2.4vw,34px);height:clamp(18px,2.4vw,34px);flex:none;color:var(--hn-deep-space)}
@keyframes mq{from{transform:translateX(0)}to{transform:translateX(-50%)}}
@media (prefers-reduced-motion:reduce){.mq-track{animation:none!important;flex-wrap:wrap;width:auto}.mq-set[aria-hidden]{display:none}.mq-row{white-space:normal}}
/* ---------- home › what is hipernova ---------- */
.mean-s{position:relative;height:340vh}
.mean-st{position:sticky;top:72px;height:calc(100vh - 72px);overflow:hidden}
.mean-st canvas{position:absolute;inset:0;width:100%;height:100%;display:block}
.mean-in{position:relative;z-index:1;height:100%;display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,360px),1fr));gap:32px;align-content:space-between;padding-top:clamp(32px,5vw,64px);padding-bottom:clamp(32px,5vw,64px);pointer-events:none}
.mean-l{display:grid;gap:20px;align-content:start;max-width:38ch}
.mean-k{color:#D7D7D7}
.mean-p{margin:0;font:400 clamp(16px,1.3vw,19px)/1.5 var(--font-body);text-wrap:pretty}
.mean-card{border-top:1px solid rgba(244,244,238,.3);padding-top:20px;margin-top:12px;display:grid;gap:12px;min-height:150px;align-content:start}
.mean-t{margin:0;font:700 clamp(22px,2vw,28px)/1.1 var(--font-display);letter-spacing:-.02em;color:var(--hn-yellow)}
.mean-d{margin:0;font:400 clamp(16px,1.3vw,18px)/1.5 var(--font-body);text-wrap:pretty}
.mean-words{list-style:none;margin:0;padding:0;align-self:end;justify-self:end;text-align:right;display:grid}
.mean-words li{font:700 clamp(28px,min(7vw,10.5vh),124px)/.92 var(--font-display);letter-spacing:-.04em;text-transform:uppercase;transition:color .6s var(--ease-field),transform .6s var(--ease-field)}
.mean-band{border-top:1px solid rgba(244,244,238,.14);padding:clamp(72px,10vw,150px) 0}
.mean-band p{margin:0;font:700 clamp(36px,5.6vw,96px)/.98 var(--font-display);letter-spacing:-.035em;text-wrap:balance}
.mean-dim{color:#D7D7D7;font-weight:500}
@media (max-width:900px){.mean-words li{font-size:clamp(24px,min(9vw,7.5vh),64px)}.mean-p{font-size:15px}.mean-card{min-height:120px}}
/* ---------- approach › method ---------- */
.method-in{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,440px),1fr));gap:clamp(56px,7vw,112px);align-items:center;padding-top:clamp(80px,11vw,160px);padding-bottom:clamp(80px,11vw,160px)}
.method-l{display:grid;gap:32px;align-content:start}
.method-k{color:#D7D7D7}
.method-h{margin:0;font:700 clamp(40px,5.4vw,88px)/.95 var(--font-display);letter-spacing:-.035em}
.method-card{border-top:1px solid rgba(244,244,238,.3);padding-top:24px;display:grid;gap:12px;min-height:180px}
.method-name{margin:0;font:700 clamp(36px,4vw,64px)/1 var(--font-display);letter-spacing:-.035em;color:var(--hn-yellow)}
.method-desc{margin:0;font:400 clamp(18px,1.6vw,22px)/1.4 var(--font-body);max-width:34ch}
.method-r{position:relative;width:min(540px,100%);aspect-ratio:1;justify-self:center;margin:24px 0}
.method-r canvas{position:absolute;inset:-6%;width:112%;height:112%;display:block;-webkit-mask-image:radial-gradient(closest-side,#000 62%,transparent 100%);mask-image:radial-gradient(closest-side,#000 62%,transparent 100%)}
.method-core{position:absolute;left:50%;top:50%;transform:translate(-50%,-50%);text-align:center;pointer-events:none;display:grid;gap:6px}
.method-core span:first-child{font:400 10px/1 var(--font-mono);letter-spacing:.14em;color:#D7D7D7}
.method-core span:last-child{font:500 11px/1 var(--font-mono);letter-spacing:.14em;color:var(--hn-yellow)}
.method-node{position:absolute;transform:translate(-50%,-50%);background:#222;color:#F4F4EE;border:1px solid #F4F4EE;border-radius:0;padding:12px 14px;cursor:pointer;display:grid;gap:4px;text-align:left;min-height:44px;transition:background .4s var(--ease-field),color .4s var(--ease-field)}
.method-node span{font:400 10px/1 var(--font-mono);letter-spacing:.14em}
.method-node b{font:700 clamp(13px,1.3vw,16px)/1 var(--font-display);letter-spacing:.02em}
.method-node.on{background:var(--hn-yellow);color:#222}
.method-node:hover:not(.on){background:#F4F4EE;color:#222}
"""
hero_js = (HERE / "src" / "hero.js").read_text(encoding="utf8")
js = (HERE / "src" / "admin.js").read_text(encoding="utf8") + "\n" + (HERE / "src" / "duo.js").read_text(encoding="utf8") + "\n" + hero_js + "\n" + (HERE / "src" / "method.js").read_text(encoding="utf8") + "\n" + (HERE / "src" / "posts.js").read_text(encoding="utf8") + "\n" + (HERE / "src" / "meaning.js").read_text(encoding="utf8") + "\n" + (HERE / "src" / "cursor.js").read_text(encoding="utf8")
import base64
FAV = "data:image/png;base64," + base64.b64encode((HERE / "favicon.png").read_bytes()).decode()
html = html.replace("</head>", '<link rel="icon" type="image/png" href="' + FAV + '">\n<link rel="apple-touch-icon" href="' + FAV + '">\n</head>', 1)
html = html.replace("</head>", config + "<style>\n" + css + "</style>\n</head>", 1)
html = html.replace("</body>", "<script>\n" + js + "\n</script>\n</body>", 1)
(HERE / "index.html").write_text(html, encoding="utf8")
print("index.html", len(html) // 1024, "KB;", len(data["projects"]), "projects")

# ---- 3. the Google Sheet seed ------------------------------------------------
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("openpyxl not installed → skipped .xlsx"); sys.exit(0)

wb = Workbook()
HEAD = PatternFill("solid", fgColor="FDEE3D")
def sheet(name, headers, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name; ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="222222"); c.fill = HEAD
    for r in rows: ws.append(r)
    for k, w in enumerate(widths): ws.column_dimensions[chr(65 + k)].width = w
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=False)
    return ws

J = lambda a: " | ".join(a)
sheet("Projects",
      ["order", "status", "slug", "title", "sub", "year", "client", "industry", "type", "role", "country", "highlight", "link", "texture", "cover", "excerpt", "body_html", "title_es", "sub_es", "excerpt_es", "body_html_es"],
      [[n + 1, p["status"], p["slug"], p["title"], p["sub"], p["year"], J(p["client"]), J(p["industry"]), J(p["type"]), J(p["role"]), J(p["country"]),
        bool(p["highlight"]), p["link"], p["texture"], p.get("cover", ""), p["excerpt"], p["html"], p["es"].get("title", ""), p["es"].get("sub", ""), "", p["es"].get("html", "")] for n, p in enumerate(data["projects"])],
      [7, 11, 34, 44, 44, 7, 24, 22, 22, 26, 16, 10, 24, 11, 40, 50, 60, 44, 44, 40, 60], first=True)
A, B, T, H = data["pages"]["approach"], data["pages"]["about"], data["pages"]["thinking"], data["pages"]["home"]
ESP = data["es"]["pages"]
sheet("Principles", ["order", "title", "text", "symbol", "title_es", "text_es"], [[n + 1, *x, *ESP["approach"]["principles"][n]] for n, x in enumerate(A["principles"])], [7, 32, 90, 16, 32, 90])
sheet("Experience", ["order", "title", "html", "title_es", "html_es"], [[n + 1, *x, *ESP["about"]["experience"][n]] for n, x in enumerate(B["experience"])], [7, 24, 120, 24, 120])
L = lambda a: "\n".join(a)
EP = data["es"]["pages"]
def _lv(v): return L(v) if isinstance(v, list) else v
PAGES = [
    ("home.hero", H["hero"], EP["home"]["hero"]), ("home.phrases", L(H["phrases"]), L(EP["home"]["phrases"])),
    ("home.lead", H["lead"], ""), ("home.bio", H["bio"], EP["home"]["bio"]),
    ("approach.lead", A["lead"], EP["approach"]["lead"]),
    ("about.intro", B["intro"], EP["about"]["intro"]), ("about.quote", B["quote"], EP["about"]["quote"]),
    ("about.education", L(B["education"]), L(EP["about"]["education"])), ("about.certifications", L(B["certifications"]), L(EP["about"]["certifications"])),
    ("about.disciplines", L(B["disciplines"]), L(EP["about"]["disciplines"])), ("about.teaching", B["teaching"], EP["about"]["teaching"]),
    ("thinking.lead", T["lead"], EP["thinking"]["lead"]), ("thinking.substack", T["substack"], EP["thinking"]["substack"]),
    ("thinking.medium", T["medium"], EP["thinking"]["medium"]), ("thinking.publications", L(T["publications"]), ""),
]
sheet("Pages", ["key", "value", "value_es"], [list(r) for r in PAGES], [26, 120, 120])
c = data["config"]
sheet("Config", ["key", "value", "value_es"], [
    ["name", c["name"], ""], ["role", c["role"], ""], ["clients", J(c["clients"]), ""], ["disciplines", J(c["disciplines"]), J(data["es"]["config"]["disciplines"])], ["substack", c["substack"], ""],
    ["linkedin", c["linkedin"], ""], ["email", c["email"], ""], ["date", c["date"], ""],
], [18, 80, 80])
sheet("Translations", ["en", "es"], [[k, v] for k, v in TRN.DICT.items()], [60, 60])
out = HERE / "database" / "hipernova-portafolio-db.xlsx"
wb.save(out); print(out.name)
